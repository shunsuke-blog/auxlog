import sys
from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
out = sys.argv[1]
bold = Font(bold=True); head = PatternFill('solid', fgColor='DDEBF7')

# A: 週ごとにシート。Day は結合セル、重量の単位なし・%混在、備考列
wb = Workbook(); wb.remove(wb.active)
days = {
 1: [('ベンチプレス', 4, '75%', 6, 'RIR2'), ('インクラインDBプレス', 3, '22', 10, ''), ('ペックフライ', 3, '40', 12, 'ラスト限界まで'), ('トライセプスプッシュダウン', 3, '25', 12, '')],
 2: [('スクワット', 4, '75%', 6, 'RIR2'), ('レッグプレス', 3, '120', 10, ''), ('レッグカール', 3, '35', 12, ''), ('カーフレイズ', 3, '自重', 20, '')],
 3: [('デッドリフト', 3, '70%', 5, ''), ('ラットプルダウン', 3, '50', 10, ''), ('シーテッドロウ', 3, '45', 12, ''), ('トレッドミル', 1, '', '20分', '傾斜5%')],
}
for w in range(1, 5):
    ws = wb.create_sheet(f'{w}週目')
    ws.append(['Day', '種目', 'セット数', '重量(kg)', '回数', '備考']); [setattr(c, 'font', bold) or setattr(c, 'fill', head) for c in ws[1]]
    r = 2
    for d, exs in days.items():
        start = r
        for name, sets, wt, reps, note in exs:
            if wt.endswith('%'): wt = f'{int(wt[:-1]) + (w-1)*2.5:g}%'
            elif wt.isdigit(): wt = str(int(wt) + (w-1)*2)
            ws.append([f'Day{d}' if r == start else None, name, sets, wt, reps, note]); r += 1
        ws.merge_cells(start_row=start, end_row=r-1, start_column=1, end_column=1)
        ws.cell(start, 1).alignment = Alignment(vertical='center')
    for col, wd in zip('ABCDEF', (8, 24, 9, 10, 8, 18)): ws.column_dimensions[col].width = wd
wb.save(f'{out}/A_週ごとのシート_結合セル.xlsx')

# B: 1枚に全部。見出し行で週・Dayを区切り、1セルに「重量×回数×セット」、略称
wb = Workbook(); ws = wb.active; ws.title = 'プログラム'
ws.append(['田中様 8週間プログラム（上半身／下半身）']); ws['A1'].font = Font(bold=True, size=14)
ws.append([])
for w in (1, 2):
    ws.append([f'■ Week {w}']); ws.cell(ws.max_row, 1).font = bold
    ws.append(['【上半身】']); 
    for row in [('BP', f'{60+w*2.5}kg×8×3', '休憩2分'), ('DBショルダー', f'{14+w}kg×10×3', ''), ('懸垂', '自重×MAX×3', '補助あり可'), ('サイドレイズ', '8kg×15×3', 'ドロップセット'), ('アームカール', '10kg×12×2', '')]:
        ws.append(['', *row])
    ws.append(['【下半身】'])
    for row in [('SQ', f'W-up 40kg×10、{80+w*5}kg×5×5', ''), ('RDL', f'{70+w*5}kg×8×3', ''), ('ブルガリアン', '各10回×3（DB10kg）', '左右'), ('ハイパーエクステンション', '15回×3', ''), ('バイク', '15分', '心拍120〜130')]:
        ws.append(['', *row])
    ws.append([])
ws.column_dimensions['A'].width = 12; ws.column_dimensions['B'].width = 24; ws.column_dimensions['C'].width = 30; ws.column_dimensions['D'].width = 16
wb.save(f'{out}/B_1枚に全部_見出し行と1セル表記.xlsx')

# C: 行が種目、列が週（横に進行）。RPE表記、セットは別列、英語名混在
wb = Workbook(); ws = wb.active; ws.title = 'Menu'
ws.append(['Day', 'Exercise', 'Sets', 'W1', 'W2', 'W3', 'W4 (deload)']); [setattr(c, 'font', bold) or setattr(c, 'fill', head) for c in ws[1]]
rows = [
 ('A', 'Back Squat', 3, '100×5 @8', '102.5×5 @8', '105×5 @8.5', '80×5 @6'),
 ('A', 'Bench Press', 3, '80×6 @8', '82.5×6 @8', '85×6 @9', '65×6 @6'),
 ('A', 'Chin-up', 3, 'BW×8', 'BW×9', 'BW×10', 'BW×6'),
 ('A', 'Plank', 3, '45秒', '50秒', '60秒', '30秒'),
 ('B', 'Deadlift', 2, '130×4 @8', '135×4 @8', '140×3 @9', '100×4 @6'),
 ('B', 'OHP', 3, '45×6', '47.5×6', '50×5', '35×6'),
 ('B', 'Barbell Row', 3, '70×8', '72.5×8', '75×8', '55×8'),
 ('B', 'Rowing machine', 1, '2000m', '2000m', '2500m', '1000m'),
]
for r in rows: ws.append(list(r))
ws.append([]); ws.append(['※ @はRPE（10=限界）。BW=自重'])
for col, wd in zip('ABCDEFG', (6, 18, 6, 12, 12, 12, 14)): ws.column_dimensions[col].width = wd
wb.save(f'{out}/C_列が週_RPEと英語名.xlsx')
